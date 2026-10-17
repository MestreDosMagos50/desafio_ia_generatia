#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
==============================================================================
MÓDULO DE COMPREENSÃO ESTRUTURADA (RF05) E GERAÇÃO DE MINUTA (RF09)
==============================================================================
- Extração estruturada pelo modelo com esquema Pydantic (saída JSON validada);
- Campos ausentes calculados pela lista de campos obrigatórios do config.json;
- Verificação de ancoragem e detecção de prompt injection (RF03);
- Minuta redigida pelo modelo a partir de dados conferidos e minimizados por
  src/privacidade.py (RF09, RF12), montada no modelo institucional e validada;
- Versão .docx com o brasão do Governo de Mato Grosso;
- Trava de segurança contra envio automático (Human-in-the-Loop).
==============================================================================
"""

import re
import json
import hashlib
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any, Tuple

from pydantic import BaseModel, Field, ValidationError, field_validator

from src.llm_client import LLMClient, LLMResponse
from src.conferencia import (
    NAO_LOCALIZADO,
    MAPA_GABARITO_SOLICITACAO,
    IGNORAR_GABARITO_SOLICITACAO,
    obter,
    eh_ausente,
    normalizar,
    carregar_gabarito,
    comparar_com_gabarito,
    verificar_ancoragem,
    detectar_injection
)
from src.privacidade import minimizar, ROTULOS_AUSENTES as ROTULOS_CAMPOS
from src.audit import nome_versionado
from src.prompts import (
    PROMPT_PRODUCAO_SISTEMA,
    PROMPT_B_ESTRUTURADO_V2,
    PROMPT_EXTRACAO_JSON_SISTEMA,
    PROMPT_EXTRACAO_JSON_USUARIO,
    VERSOES_PROMPTS
)

FORMATO_DATA = "Data no formato DD/MM/AAAA, ou 'não localizado'."

# Menções que, se ausentes das fontes, indicam possível informação inventada (texto normalizado)
PADROES_INVENCAO = (
    ("lei ou norma", r"\b(lei|decreto|portaria|resolucao|instrucao normativa)\b"),
    ("edital", r"\bedita(l|is|licia|licias|licio)\b"),
    ("artigo de norma", r"\bart(igo)?\.?\s*\d"),
    ("prazo", r"\b(prazo|prazos|dias uteis|improrrogavel)\b"),
    ("documento específico", r"\b(certidao|historico escolar|comprovante de residencia|foto 3x4|titulo de eleitor|reservista|carteira de vacinacao|certificado de conclusao)\b"),
)
PADRAO_TERMO_DECISAO = r"\b(aprovad[ao]s?|deferid[ao]s?|homologad[ao]s?|matricula (esta )?(efetivada|confirmada))\b"


def _vazio_para_nao_localizado(valor):
    """Campo recebido vazio ('' ou só espaços) vira 'não localizado' (RF05)."""
    if isinstance(valor, str) and not valor.strip():
        return NAO_LOCALIZADO
    return valor


class SolicitanteSchema(BaseModel):
    nome_completo: str = Field(default=NAO_LOCALIZADO, description="Nome completo do solicitante")
    data_nascimento: str = Field(default=NAO_LOCALIZADO, description=FORMATO_DATA)
    naturalidade: str = Field(default=NAO_LOCALIZADO, description="Cidade de nascimento")
    cpf: str = Field(default=NAO_LOCALIZADO)
    rg: str = Field(default=NAO_LOCALIZADO)
    orgao_expedidor: str = Field(default=NAO_LOCALIZADO, description="Órgão expedidor do RG")
    uf: str = Field(default=NAO_LOCALIZADO, description="UF de expedição do RG")
    nome_mae: str = Field(default=NAO_LOCALIZADO)
    nome_pai: str = Field(default=NAO_LOCALIZADO)
    email: str = Field(default=NAO_LOCALIZADO)
    telefone_celular: str = Field(default=NAO_LOCALIZADO)
    telefone_residencial: str = Field(default=NAO_LOCALIZADO)

    _normalizar_vazios = field_validator("*", mode="before")(_vazio_para_nao_localizado)


class CursoSchema(BaseModel):
    nome: str = Field(default=NAO_LOCALIZADO, description="Nome do curso")
    eixo_tecnologico: str = Field(default=NAO_LOCALIZADO)
    carga_horaria: str = Field(default=NAO_LOCALIZADO)
    turno: str = Field(default=NAO_LOCALIZADO)
    data_inicio: str = Field(default=NAO_LOCALIZADO, description=FORMATO_DATA)
    data_termino: str = Field(default=NAO_LOCALIZADO, description=FORMATO_DATA)
    turma: str = Field(default=NAO_LOCALIZADO)

    _normalizar_vazios = field_validator("*", mode="before")(_vazio_para_nao_localizado)


class ExtracaoSolicitacao(BaseModel):
    """Campos preenchidos pelo modelo: é o esquema JSON enviado ao provedor."""
    protocolo: str = Field(default=NAO_LOCALIZADO)
    data_solicitacao: str = Field(default=NAO_LOCALIZADO, description=FORMATO_DATA)
    canal_origem: str = Field(default=NAO_LOCALIZADO, description="Canal informado no documento (ex.: E-mail)")
    solicitante: SolicitanteSchema = Field(default_factory=SolicitanteSchema)
    curso: CursoSchema = Field(default_factory=CursoSchema)
    assunto: str = Field(default=NAO_LOCALIZADO, description="Assunto em uma frase curta")
    pedido_detalhado: str = Field(default=NAO_LOCALIZADO, description="O que o solicitante pede, incluindo as perguntas feitas")
    declarou_ciencia_faltas: bool = Field(default=False)
    texto_ciencia_faltas: str = Field(default=NAO_LOCALIZADO, description="Trecho literal da declaração de ciência das regras de faltas")
    documentos_anexados: List[str] = Field(default_factory=list)
    categoria_tema: str = Field(default="OUTROS")
    urgencia: str = Field(default="MEDIA", description="BAIXA, MEDIA ou ALTA")
    nivel_confianca: str = Field(default="BAIXO", description="ALTO, MEDIO ou BAIXO")
    suspeita_prompt_injection: bool = Field(default=False)
    trechos_suspeitos: List[str] = Field(default_factory=list)

    _normalizar_vazios = field_validator(
        "protocolo", "data_solicitacao", "canal_origem", "assunto", "pedido_detalhado", "texto_ciencia_faltas",
        mode="before"
    )(_vazio_para_nao_localizado)


class SolicitacaoEstruturada(ExtracaoSolicitacao):
    """Registro completo da solicitação (RF05): extração, lacunas, alertas e rastreabilidade."""
    campos_ausentes: List[str] = Field(default_factory=list)
    correcoes_revisor: List[str] = Field(default_factory=list, description="Valores corrigidos na conferência humana")
    necessita_revisao_humana: bool = Field(default=True)
    alertas: List[str] = Field(default_factory=list)
    texto_original: str = Field(default="")
    arquivo_origem: str = Field(default="")
    hash_origem_sha256: str = Field(default="")


class MinutaGerada(BaseModel):
    """Conteúdo redigido pelo modelo; cabeçalho, lacunas, fontes e status são inseridos pelo código."""
    assunto: str
    saudacao: str
    paragrafos: List[str]
    fechamento: str
    alertas_para_revisor: List[str] = Field(default_factory=list)


class GeradorMinuta:
    """Extração estruturada e geração da minuta institucional com bloqueio de envio automático."""

    def __init__(self, config_path: str = "config.json", llm_client: Optional[LLMClient] = None):
        self.raiz_projeto = Path(__file__).resolve().parent.parent
        caminho_config = self.raiz_projeto / config_path
        self.config = json.loads(caminho_config.read_text(encoding="utf-8")) if caminho_config.exists() else {}
        self.llm_client = llm_client or LLMClient(config_path)
        self.categorias = self.config.get("classificacao", {}).get("categorias_validas", ["OUTROS"])
        self.campos_obrigatorios = self.config.get("extracao", {}).get("campos_obrigatorios", [])
        self.dir_processados = self.raiz_projeto / "dados" / "processados"
        self.dir_minutas = self.raiz_projeto / "resultados" / "minutas"

    @staticmethod
    def calcular_sha256(caminho_arquivo: Path) -> str:
        """Gera hash SHA-256 para rastreabilidade estrita da fonte de dados."""
        if not caminho_arquivo.exists():
            return "HASH_NAO_DISPONIVEL"
        h = hashlib.sha256()
        with open(caminho_arquivo, "rb") as f:
            while chunk := f.read(8192):
                h.update(chunk)
        return h.hexdigest()

    def _relativo(self, caminho: Path) -> str:
        try:
            return caminho.resolve().relative_to(self.raiz_projeto).as_posix()
        except ValueError:
            return caminho.as_posix()

    @staticmethod
    def nome_base(arquivo: str) -> str:
        """'dados/originais/entrada_texto.txt' -> 'texto' (usado nos nomes dos artefatos)."""
        stem = Path(arquivo).stem
        return stem[len("entrada_"):] if stem.startswith("entrada_") else stem

    def calcular_campos_ausentes(self, registro: Dict[str, Any]) -> List[str]:
        """Lista os campos obrigatórios (config.json) sem valor no registro."""
        ausentes = []
        for caminho in self.campos_obrigatorios:
            if caminho == "documentos_comprobatorios_anexos":
                if not registro.get("documentos_anexados"):
                    ausentes.append(caminho)
            elif eh_ausente(obter(registro, caminho)):
                ausentes.append(caminho.split(".")[-1])
        return ausentes

    def extrair_dados(
        self,
        texto: str,
        canal: str,
        arquivo_origem: Path,
        salvar_como: Optional[str] = None
    ) -> Tuple[Optional[SolicitacaoEstruturada], LLMResponse, str]:
        """
        Extração estruturada de uma solicitação (RF05).
        Retorna (registro, resposta do modelo, prompt enviado). Se o provedor falhar, registro é None.
        """
        prompt_sistema = PROMPT_EXTRACAO_JSON_SISTEMA.format(categorias=", ".join(self.categorias))
        prompt_usuario = PROMPT_EXTRACAO_JSON_USUARIO.format(canal=canal, conteudo_documento=texto)
        prompt_completo = f"{prompt_sistema}\n{prompt_usuario}"

        resposta = self.llm_client.gerar(prompt_usuario, system_instruction=prompt_sistema, schema=ExtracaoSolicitacao)
        if not resposta.sucesso:
            return None, resposta, prompt_completo

        alertas = []
        try:
            extracao = ExtracaoSolicitacao.model_validate_json(resposta.conteudo)
        except ValidationError as e:
            # Critério de aceite do RF05: o JSON continua válido mesmo sem dados aproveitáveis
            extracao = ExtracaoSolicitacao()
            alertas.append(f"Resposta do modelo fora do esquema ({e.error_count()} erro(s)); campos marcados como não localizados.")

        dados = extracao.model_dump()
        if dados["categoria_tema"] not in self.categorias:
            alertas.append(f"Categoria '{dados['categoria_tema']}' fora da lista documentada; reclassificada como OUTROS.")
            dados["categoria_tema"] = "OUTROS"
        if eh_ausente(dados["canal_origem"]):
            dados["canal_origem"] = canal

        # Detecção heurística complementar à sinalização do modelo (RF03)
        trechos = detectar_injection(texto)
        if trechos and not dados["suspeita_prompt_injection"]:
            alertas.append("A detecção heurística encontrou instruções dirigidas à IA que o modelo não sinalizou.")
        if trechos:
            dados["suspeita_prompt_injection"] = True
            dados["trechos_suspeitos"] = list(dict.fromkeys(dados["trechos_suspeitos"] + trechos))
        if dados["suspeita_prompt_injection"]:
            alertas.append("Suspeita de prompt injection: a solicitação contém instruções dirigidas ao sistema.")

        alertas.extend(verificar_ancoragem(dados, texto))

        registro = SolicitacaoEstruturada(
            **dados,
            campos_ausentes=self.calcular_campos_ausentes(dados),
            necessita_revisao_humana=True,
            alertas=alertas,
            texto_original=texto,
            arquivo_origem=self._relativo(arquivo_origem),
            hash_origem_sha256=self.calcular_sha256(arquivo_origem)
        )

        if salvar_como:
            self.dir_processados.mkdir(parents=True, exist_ok=True)
            (self.dir_processados / salvar_como).write_text(registro.model_dump_json(indent=2), encoding="utf-8")

        return registro, resposta, prompt_completo

    def extrair_dados_texto(self, arquivo_texto: str) -> Tuple[Optional[SolicitacaoEstruturada], LLMResponse, str, Path]:
        """Lê a solicitação textual (arquivo original inalterado) e extrai a estrutura (RF05)."""
        caminho = self.raiz_projeto / arquivo_texto
        texto_bruto = caminho.read_text(encoding="utf-8")
        saida = self.dir_processados / f"extracao_{self.nome_base(arquivo_texto)}.json"
        registro, resposta, prompt = self.extrair_dados(texto_bruto, "texto", caminho, salvar_como=saida.name)
        return registro, resposta, prompt, saida

    def comparar_com_gabarito(self, registro: SolicitacaoEstruturada, arquivo_entrada: str) -> Optional[Dict[str, Any]]:
        """Compara a extração com o gabarito elaborado pela equipe (RF05 / RF13) e salva o resultado."""
        gabarito = carregar_gabarito(self.raiz_projeto, self.config, arquivo_entrada)
        if not gabarito:
            return None
        comparacao = comparar_com_gabarito(
            registro.model_dump(),
            gabarito.get("dados_extraidos_esperados", {}),
            mapa=MAPA_GABARITO_SOLICITACAO,
            ignorar=IGNORAR_GABARITO_SOLICITACAO
        )
        comparacao["caso_id"] = gabarito.get("caso_id")
        saida = self.dir_processados / f"comparacao_{self.nome_base(arquivo_entrada)}_gabarito.json"
        saida.write_text(json.dumps(comparacao, indent=2, ensure_ascii=False), encoding="utf-8")
        return comparacao

    def gerar_minuta_institucional(self, registro: SolicitacaoEstruturada, fontes: List[Dict[str, str]]) -> Dict[str, Any]:
        """
        Produz a minuta institucional (RF09) a partir do registro conferido pelo operador.
        `fontes`: arquivos de origem usados ({'canal', 'arquivo', 'sha256', 'texto'}).
        """
        dados = minimizar(registro.model_dump())
        prompt_usuario = PROMPT_B_ESTRUTURADO_V2.format(dados_json=json.dumps(dados, ensure_ascii=False, indent=2))
        resultado: Dict[str, Any] = {
            "sucesso": False,
            "protocolo": registro.protocolo,
            "prompt_id": "PROMPT_PRODUCAO_SISTEMA+PROMPT_B_ESTRUTURADO_V2",
            "prompt_versao": f"{VERSOES_PROMPTS['PROMPT_PRODUCAO_SISTEMA']}+{VERSOES_PROMPTS['PROMPT_B_ESTRUTURADO_V2']}",
            "prompt_completo": f"{PROMPT_PRODUCAO_SISTEMA}\n{prompt_usuario}",
            "status": "NAO_GERADA",
            "envio_automatico_bloqueado": True
        }

        resposta = self.llm_client.gerar(prompt_usuario, system_instruction=PROMPT_PRODUCAO_SISTEMA, schema=MinutaGerada)
        resultado["resposta"] = resposta
        if not resposta.sucesso:
            resultado["erro"] = resposta.erro
            return resultado
        try:
            conteudo = MinutaGerada.model_validate_json(resposta.conteudo)
        except ValidationError as e:
            resposta.marcar_falha("RESPOSTA_FORA_DO_ESQUEMA", f"A minuta retornada não segue o esquema esperado ({e.error_count()} erro(s)).")
            resultado["erro"] = resposta.erro
            return resultado

        texto_fontes = "\n".join(f.get("texto", "") for f in fontes) or registro.texto_original
        alertas = self.validar_minuta(conteudo, registro, texto_fontes)

        # Cada geração grava arquivos novos (protocolo + data e hora): versões anteriores são preservadas
        protocolo_arquivo = re.sub(r"[^\w.-]", "_", registro.protocolo)
        self.dir_minutas.mkdir(parents=True, exist_ok=True)
        nome = nome_versionado(self.dir_minutas, f"minuta_institucional_{protocolo_arquivo}", [".txt", ".docx", "_validacao.json"])
        caminho_txt = self.dir_minutas / f"{nome}.txt"
        caminho_docx = self.dir_minutas / f"{nome}.docx"
        caminho_validacao = self.dir_minutas / f"{nome}_validacao.json"

        corpo = self._renderizar_txt(registro, conteudo, fontes)
        caminho_txt.write_text(corpo, encoding="utf-8")
        erro_docx = self._salvar_docx(caminho_docx, registro, conteudo, fontes)
        if erro_docx:
            alertas.append(erro_docx)
        caminho_validacao.write_text(json.dumps({
            "protocolo": registro.protocolo,
            "gerada_em": datetime.now().isoformat(),
            "alertas_validacao_automatica": alertas,
            "dados_enviados_ao_modelo": dados
        }, indent=2, ensure_ascii=False), encoding="utf-8")

        resultado.update({
            "sucesso": True,
            "caminho_arquivo": str(caminho_txt),
            "caminho_docx": str(caminho_docx) if not erro_docx else None,
            "caminho_validacao": str(caminho_validacao),
            "conteudo_minuta": corpo,
            "conteudo_gerado": conteudo.model_dump(),
            "alertas": alertas,
            "status": "PENDENTE_REVISAO_HUMANA"
        })
        return resultado

    def validar_minuta(self, conteudo: MinutaGerada, registro: SolicitacaoEstruturada, texto_fontes: str) -> List[str]:
        """Verificações automáticas antes da revisão humana (RF03 / RF09). Não substituem o revisor."""
        texto = "\n".join([conteudo.assunto, conteudo.saudacao, *conteudo.paragrafos, conteudo.fechamento])
        texto_normalizado = normalizar(texto)
        fontes_normalizadas = normalizar(texto_fontes)

        alertas = [f"Alerta do modelo: {a}" for a in conteudo.alertas_para_revisor]
        if re.search(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", texto):
            alertas.append("A minuta contém um número no formato de CPF.")
        if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", texto):
            alertas.append("A minuta reproduz um endereço de e-mail.")
        if "***" in texto:
            alertas.append("A minuta reproduz um identificador mascarado (CPF, RG ou e-mail).")
        for rotulo, padrao in PADROES_INVENCAO:
            if re.search(padrao, texto_normalizado) and not re.search(padrao, fontes_normalizadas):
                alertas.append(f"A minuta menciona {rotulo}, que não consta nas fontes: possível informação inventada.")
        if re.search(PADRAO_TERMO_DECISAO, texto_normalizado):
            alertas.append("A minuta usa termo de decisão (aprovada/deferida/homologada): confirmar que não comunica decisão.")
        if registro.declarou_ciencia_faltas and not eh_ausente(registro.texto_ciencia_faltas):
            if normalizar(registro.texto_ciencia_faltas) not in texto_normalizado:
                alertas.append("A declaração de ciência de faltas não foi citada literalmente: conferir se o sentido foi preservado.")
        if registro.suspeita_prompt_injection:
            alertas.append("A solicitação contém suspeita de prompt injection: revisar com atenção.")
        return alertas

    def _renderizar_txt(self, registro: SolicitacaoEstruturada, conteudo: MinutaGerada, fontes: List[Dict[str, str]]) -> str:
        """Monta a minuta no modelo institucional definido pela equipe."""
        separador = "=" * 80
        linhas = [
            separador,
            "GOVERNO DO ESTADO DE MATO GROSSO",
            "SECRETARIA DE ESTADO DE CIÊNCIA, TECNOLOGIA E INOVAÇÃO - SECITECI/MT",
            "ESCOLA TÉCNICA ESTADUAL DE CUIABÁ",
            separador,
            "MINUTA DE RESPOSTA AO CIDADÃO",
            f"Protocolo: {registro.protocolo}",
            f"Data da solicitação: {registro.data_solicitacao} | Canal: {registro.canal_origem}",
            f"Data de geração da minuta: {datetime.now().strftime('%d/%m/%Y')}",
            f"Assunto: {conteudo.assunto}",
            separador,
            "",
            conteudo.saudacao,
            ""
        ]
        for paragrafo in conteudo.paragrafos:
            linhas += [paragrafo, ""]
        if registro.campos_ausentes:
            linhas.append("INFORMAÇÕES NÃO LOCALIZADAS NOS DADOS RECEBIDOS:")
            linhas += [f"- {ROTULOS_CAMPOS.get(c, c)}: {NAO_LOCALIZADO}" for c in registro.campos_ausentes]
            linhas.append("")
        linhas += [
            conteudo.fechamento,
            "",
            "Secretaria Escolar - Escola Técnica Estadual de Cuiabá",
            "",
            separador,
            "FONTES UTILIZADAS (RF09 / RF11):"
        ]
        linhas += [f"- {f['arquivo']} ({f['canal']}) - SHA-256: {f['sha256']}" for f in fontes]
        linhas += [
            separador,
            "[STATUS: MINUTA GERADA COM APOIO DE IA - PENDENTE DE REVISÃO HUMANA]",
            "O ENVIO AUTOMÁTICO DESTA RESPOSTA ESTÁ BLOQUEADO. A DECISÃO CABE AO SERVIDOR RESPONSÁVEL.",
            separador,
            ""
        ]
        return "\n".join(linhas)

    def _salvar_docx(self, caminho: Path, registro: SolicitacaoEstruturada, conteudo: MinutaGerada, fontes: List[Dict[str, str]]) -> Optional[str]:
        """Gera a versão .docx com o brasão no cabeçalho. Retorna mensagem de erro ou None."""
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm, Pt, RGBColor
        except ImportError:
            return "Versão .docx não gerada: biblioteca python-docx não instalada (pip install -r requirements.txt)."

        documento = Document()
        cabecalho = documento.sections[0].header
        paragrafo = cabecalho.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brasao = self.raiz_projeto / self.config.get("arquivos_originais", {}).get("brasao", "dados/originais/brasao.png")
        if brasao.exists():
            paragrafo.add_run().add_picture(str(brasao), width=Cm(2.0))
        for linha in ("GOVERNO DO ESTADO DE MATO GROSSO",
                      "SECRETARIA DE ESTADO DE CIÊNCIA, TECNOLOGIA E INOVAÇÃO - SECITECI/MT",
                      "ESCOLA TÉCNICA ESTADUAL DE CUIABÁ"):
            p = cabecalho.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(linha)
            run.bold = True
            run.font.size = Pt(10)

        documento.add_heading("Minuta de resposta ao cidadão", level=1)
        documento.add_paragraph(
            f"Protocolo: {registro.protocolo}\n"
            f"Data da solicitação: {registro.data_solicitacao} | Canal: {registro.canal_origem}\n"
            f"Data de geração da minuta: {datetime.now().strftime('%d/%m/%Y')}\n"
            f"Assunto: {conteudo.assunto}"
        )
        documento.add_paragraph(conteudo.saudacao)
        for texto in conteudo.paragrafos:
            documento.add_paragraph(texto)
        if registro.campos_ausentes:
            documento.add_paragraph("Informações não localizadas nos dados recebidos:").runs[0].bold = True
            for campo in registro.campos_ausentes:
                documento.add_paragraph(f"{ROTULOS_CAMPOS.get(campo, campo)}: {NAO_LOCALIZADO}", style="List Bullet")
        documento.add_paragraph(conteudo.fechamento)
        documento.add_paragraph("Secretaria Escolar - Escola Técnica Estadual de Cuiabá")

        fontes_par = documento.add_paragraph()
        fontes_run = fontes_par.add_run(
            "Fontes utilizadas: " + "; ".join(f"{f['arquivo']} ({f['canal']}, SHA-256 {f['sha256'][:16]}...)" for f in fontes)
        )
        fontes_run.font.size = Pt(8)
        status = documento.add_paragraph().add_run(
            "MINUTA GERADA COM APOIO DE IA - PENDENTE DE REVISÃO HUMANA. "
            "O envio automático está bloqueado; a decisão cabe ao servidor responsável."
        )
        status.bold = True
        status.font.color.rgb = RGBColor(0xB4, 0x28, 0x28)
        documento.save(str(caminho))
        return None
